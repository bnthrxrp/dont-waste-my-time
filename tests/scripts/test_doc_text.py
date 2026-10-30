"""Tests for doc_text.py: extraction from the real fixtures plus small synthetic files."""
import json
import tempfile
import unittest
from pathlib import Path

import _helpers as h
import doc_text


def by_location(units):
    return {u["location"]: u for u in units}


class FixtureExtraction(unittest.TestCase):
    def test_weak_deck_titles_and_filler_slides(self):
        titles = doc_text.slide_titles(h.WEAK_DECK)
        self.assertEqual(len(titles), 10)
        self.assertEqual(titles[0], {"slide": 1, "title": "Q3 Business Review"})
        self.assertEqual([t["title"] for t in titles[1:4]], ["Agenda", "Background", "Sales Overview"])
        self.assertEqual(titles[-1]["title"], "Thank You — Questions?")

    def test_weak_deck_3d_pie_chart_does_not_crash_and_is_read_from_xml(self):
        units, warnings = doc_text.extract_with_warnings(h.WEAK_DECK)
        chart = by_location(units)["slide 5 / chart"]
        self.assertEqual(chart["kind"], "chart")
        for expected in ("North America", "Oceania", "Q3 revenue ($M)", "31.4", "1.5"):
            self.assertIn(expected, chart["text"])
        self.assertNotIn("unparsed", chart["text"])
        self.assertEqual(warnings, [])

    def test_chart_values_are_flagged_as_machine_numbers(self):
        chart = by_location(doc_text.extract_units(h.WEAK_DECK))["slide 5 / chart"]
        spans = chart["machine_spans"]
        self.assertTrue(spans)
        start, end = spans[0]
        self.assertEqual(chart["text"][start:end], "31.4 | 22.8 | 9.7 | 7.9 | 5.6 | 4.1 | 3.2 | 1.5")

    def test_weak_deck_table_bullets_and_kinds(self):
        units = by_location(doc_text.extract_units(h.WEAK_DECK))
        self.assertEqual(units["slide 6 / table r1c2"]["text"], "Rev Q3 ($M)")
        self.assertEqual(units["slide 6 / table r3c2"]["text"], "22.8")
        self.assertEqual(units["slide 6 / table r3c2"]["kind"], "table")
        body = units["slide 4 / body"]
        self.assertEqual(body["kind"], "body")
        self.assertEqual(len(body["paragraphs"]), 10)
        self.assertEqual(body["text"].split("\n")[0], "Revenue was $86.2M in Q3")
        self.assertEqual(units["slide 4 / title"]["kind"], "title")

    def test_pt_deck_chart_table_and_notes(self):
        units = by_location(doc_text.extract_units(h.PT_DECK))
        self.assertEqual(units["slide 3 / table r2c2"]["text"], "21,4")
        self.assertEqual(units["slide 3 / table r6c1"]["text"], "Total")
        chart = units["slide 4 / chart"]["text"]
        self.assertIn("Categories: Sudeste | Sul | Centro-Oeste | Nordeste | Norte", chart)
        self.assertIn('Series "Positivação 3T26 (%)": 74.8 | 72.1 | 69 | 66.4 | 61.7', chart)
        self.assertEqual(units["slide 1 / notes"]["kind"], "notes")
        self.assertIn("board global", units["slide 1 / notes"]["text"])
        self.assertIn("Positivação = % dos PDVs", units["slide 4 / notes"]["text"])
        self.assertIn("IQVIA", units["slide 5 / body 2"]["text"])  # slide 5 has two text boxes

    def test_markdown_keeps_table_rows_and_headings(self):
        units = doc_text.extract_units(h.REPORT_MD)
        locs = by_location(units)
        self.assertEqual(locs["line 4"]["kind"], "heading")
        self.assertEqual(locs["line 4"]["text"], "1. Introduction")
        row = locs["line 14"]
        self.assertEqual(row["kind"], "line")
        self.assertTrue(row["text"].startswith("| Northeast | 78.4 |"))
        self.assertFalse(any(set(u["text"]) <= set("|-: ") for u in units), "table separator rows are noise")

    def test_xlsx_cells_with_numbers(self):
        units = by_location(doc_text.extract_units(h.PRICING_XLSX))
        self.assertEqual(units["SKUs!A1"]["text"], "sku")
        self.assertEqual(units["SKUs!A2"]["text"], "REL-24")
        self.assertEqual(units["SKUs!D2"]["text"], "8.99")
        self.assertEqual(units["SKUs!G2"]["text"], "-0.6")
        self.assertEqual(units["SKUs!D2"]["machine_spans"], [[0, 4]])
        self.assertIn("Assumptions!B2", units)
        self.assertEqual(units["Assumptions!B2"]["text"], "5.5%")
        self.assertNotIn("machine_spans", units["Assumptions!B2"])  # text cell, not a numeric one

    def test_csv_fixture_rows(self):
        units = doc_text.extract_units(h.KPI_CSV)
        self.assertEqual(units[0]["location"], "row 1")
        self.assertTrue(units[0]["text"].startswith("month | channel | net_sales_musd"))
        self.assertTrue(units[1]["text"].startswith("2026-01 | Retail chains | 22.26"))


class SyntheticFiles(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_pptx_group_shapes_are_walked_recursively(self):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Grouped content"
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text_frame.text = "Inside the group"
        nested = group.shapes.add_group_shape()
        nested.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1)).text_frame.text = "Nested group text"
        path = self.tmp / "g.pptx"
        prs.save(path)
        text = " ".join(u["text"] for u in doc_text.extract_units(path))
        self.assertIn("Inside the group", text)
        self.assertIn("Nested group text", text)

    def test_docx_paragraphs_headings_and_tables(self):
        import docx
        document = docx.Document()
        document.add_heading("Report title", 0)
        document.add_heading("Section", 1)
        document.add_paragraph("Normal paragraph with 5.2% growth.")
        table = document.add_table(rows=2, cols=3)
        for c, label in enumerate("ABC"):
            table.cell(0, c).text = label
        table.cell(1, 0).text = "1"
        table.cell(1, 1).merge(table.cell(1, 2)).text = "merged"
        path = self.tmp / "d.docx"
        document.save(path)
        units = by_location(doc_text.extract_units(path))
        self.assertEqual((units["paragraph 1"]["kind"], units["paragraph 2"]["kind"], units["paragraph 3"]["kind"]),
                         ("title", "heading", "body"))
        self.assertEqual(units["table 1 r2c2"]["text"], "merged")
        self.assertNotIn("table 1 r2c3", units)  # merged cell is reported once

    def test_html_visible_text_only_with_attributes(self):
        html = self.tmp / "p.html"
        html.write_text(
            "<!doctype html><html><head><title>Q3 &amp; Results</title><style>.x{color:red}</style>"
            "<script>var secret='HIDDEN_SCRIPT_TEXT';</script></head><body>"
            "<h1>Revenue grew <em>6.2%</em></h1><p>Para <a title=\"Source: IQVIA\" href=#>link</a> text"
            "<div hidden>HIDDEN_ATTR</div><div style=\"display:none\">HIDDEN_STYLE</div>"
            "<button aria-label=\"Close dialog\">X</button><noscript>NOSCRIPT_TEXT</noscript>"
            "<table><tr><th>Region</th><th>Rev</th></tr><tr><td>NA</td><td>31,4</td></tr></table></body></html>",
            encoding="utf-8")
        units = doc_text.extract_units(html)
        joined = " | ".join(u["text"] for u in units)
        for hidden in ("HIDDEN_SCRIPT_TEXT", "HIDDEN_ATTR", "HIDDEN_STYLE", "NOSCRIPT_TEXT", "color:red"):
            self.assertNotIn(hidden, joined)
        locs = by_location(units)
        self.assertEqual(locs["html / title"]["text"], "Q3 & Results")
        heading = next(u for u in units if u["kind"] == "heading")
        self.assertEqual(heading["text"], "Revenue grew 6.2%")
        self.assertIn("Source: IQVIA", joined)
        self.assertIn("Close dialog", joined)
        self.assertEqual(locs["html / table 1 r2c2"]["text"], "31,4")
        order = [u["text"] for u in units]
        self.assertLess(order.index("Revenue grew 6.2%"), order.index("Close dialog"))

    def test_csv_semicolon_delimiter_and_cp1252_encoding(self):
        path = self.tmp / "br.csv"
        path.write_bytes("mês;valor\njaneiro;1.234,5\n".encode("cp1252"))
        units = doc_text.extract_units(path)
        self.assertEqual(units[0]["text"], "mês | valor")
        self.assertEqual(units[1]["text"], "janeiro | 1.234,5")

    def test_missing_library_gives_an_install_hint(self):
        import sys
        saved = sys.modules.get("pptx", "absent")
        sys.modules["pptx"] = None  # makes `import pptx` raise ImportError, as on a machine without python-pptx
        try:
            with self.assertRaises(doc_text.MissingDependency) as caught:
                doc_text.extract_units(h.WEAK_DECK)
        finally:
            if saved == "absent":
                del sys.modules["pptx"]
            else:
                sys.modules["pptx"] = saved
        self.assertIn("pip install python-pptx", str(caught.exception))

    def test_language_detection(self):
        self.assertEqual(doc_text.detect_language("O faturamento cresceu em relação ao ano anterior e a margem caiu"), "pt")
        self.assertEqual(doc_text.detect_language("Revenue grew compared with the prior year and margin fell"), "en")
        self.assertEqual(doc_text.detect_language(""), "en")


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def build_rich_docx(path: Path) -> None:
    """A minimal hand-written .docx with footnotes, endnotes, comments, header/footer and a text box.

    Hand-written because python-docx cannot create footnotes; the file is still opened by python-docx
    (for the body) and by the zip reader (for the extra parts), exactly like a real document.
    """
    import zipfile

    def note_part(root, tag, items):
        entries = "".join(
            f'<w:{tag} w:id="{i}"{extra}><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:{tag}>' for i, extra, text in items)
        return f'<?xml version="1.0" encoding="UTF-8"?><w:{root} xmlns:w="{W_NS}">{entries}</w:{root}>'

    # Footnote ids 0/1 are Word's separators; ids 2 and 3 are real. The body references 3 first, so 3 is "footnote 1".
    footnotes = note_part("footnotes", "footnote", [
        (0, ' w:type="separator"', "SEPARATOR"), (1, ' w:type="continuationSeparator"', "CONT"),
        (2, "", "Second reference: IQVIA data to August"), (3, "", "First reference: source is finance")])
    endnotes = note_part("endnotes", "endnote", [(2, "", "Endnote about methodology 2026")])
    comments = (f'<?xml version="1.0" encoding="UTF-8"?><w:comments xmlns:w="{W_NS}">'
                f'<w:comment w:id="0" w:author="Ana Reviewer"><w:p><w:r><w:t>Check the 12% figure</w:t></w:r></w:p></w:comment>'
                f'</w:comments>')
    header = (f'<?xml version="1.0" encoding="UTF-8"?><w:hdr xmlns:w="{W_NS}"><w:p><w:r><w:t>Confidential - Vitalis</w:t></w:r></w:p></w:hdr>')
    footer = (f'<?xml version="1.0" encoding="UTF-8"?><w:ftr xmlns:w="{W_NS}"><w:p><w:r><w:t>3</w:t></w:r></w:p></w:ftr>')
    document = (
        f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{W_NS}" '
        f'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:r="{REL}"><w:body>'
        '<w:p><w:r><w:t>Body paragraph one</w:t></w:r><w:r><w:footnoteReference w:id="3"/></w:r>'
        '<w:r><w:t xml:space="preserve"> and more </w:t></w:r><w:r><w:footnoteReference w:id="2"/></w:r>'
        '<w:r><w:endnoteReference w:id="2"/></w:r>'
        '<w:commentRangeStart w:id="0"/><w:r><w:t>commented words</w:t></w:r><w:r><w:commentReference w:id="0"/></w:r></w:p>'
        '<w:p><w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wps:txbx '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><w:txbxContent><w:p><w:r>'
        '<w:t>Callout: margin 41,2%</w:t></w:r></w:p></w:txbxContent></wps:txbx></w:drawing></mc:Choice>'
        '<mc:Fallback><w:pict><w:txbxContent><w:p><w:r><w:t>Callout: margin 41,2%</w:t></w:r></w:p></w:txbxContent>'
        '</w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>'
        '<w:sectPr><w:headerReference w:type="default" r:id="rIdH"/><w:footerReference w:type="default" r:id="rIdF"/></w:sectPr>'
        '</w:body></w:document>')
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '</Types>')
    root_rels = (f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                 f'<Relationship Id="rId1" Type="{REL}/officeDocument" Target="word/document.xml"/></Relationships>')
    doc_rels = (f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                f'<Relationship Id="rIdH" Type="{REL}/header" Target="header1.xml"/>'
                f'<Relationship Id="rIdF" Type="{REL}/footer" Target="footer1.xml"/></Relationships>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", root_rels)
        z.writestr("word/document.xml", document)
        z.writestr("word/_rels/document.xml.rels", doc_rels)
        for name, xml in (("footnotes", footnotes), ("endnotes", endnotes), ("comments", comments),
                          ("header1", header), ("footer1", footer)):
            z.writestr(f"word/{name}.xml", xml)


class ExtraParts(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_docx_footnotes_endnotes_comments_headers_and_textboxes(self):
        path = self.tmp / "rich.docx"
        build_rich_docx(path)
        units = doc_text.extract_units(path)
        locs = by_location(units)
        self.assertEqual(locs["paragraph 1"]["text"], "Body paragraph one and more commented words")
        # Numbered in the order the body references them, not by w:id; separators are skipped.
        self.assertEqual(locs["footnote 1"]["text"], "First reference: source is finance")
        self.assertEqual(locs["footnote 2"]["text"], "Second reference: IQVIA data to August")
        self.assertEqual(locs["footnote 1"]["kind"], "notes")
        self.assertNotIn("footnote 3", locs)
        self.assertEqual(locs["endnote 1"]["text"], "Endnote about methodology 2026")
        self.assertEqual(locs["comment 1"]["text"], "Check the 12% figure")
        self.assertEqual(locs["comment 1"]["author"], "Ana Reviewer")
        self.assertEqual(locs["header"]["text"], "Confidential - Vitalis")
        self.assertNotIn("footer", locs, "a footer that holds only a page number is not content")
        # The text box is stored twice (DrawingML + VML fallback) but reported once, and not inside the paragraph.
        textboxes = [u for u in units if u["location"].startswith("textbox")]
        self.assertEqual([u["text"] for u in textboxes], ["Callout: margin 41,2%"])
        self.assertEqual(sum("Callout" in u["text"] for u in units), 1)
        self.assertNotIn("SEPARATOR", " ".join(u["text"] for u in units))

    def test_docx_without_extra_parts_is_unaffected(self):
        import docx
        document = docx.Document()
        document.add_paragraph("Only a body")
        path = self.tmp / "plain.docx"
        document.save(path)
        self.assertEqual([u["location"] for u in doc_text.extract_units(path)], ["paragraph 1"])

    def _deck_with_alt_and_template_text(self):
        import struct
        import zlib
        from pptx import Presentation
        from pptx.oxml.shapes.autoshape import CT_Shape
        from pptx.util import Inches

        def png_1x1() -> bytes:
            def chunk(tag, data):
                body = tag + data
                return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))
            raw = zlib.compress(b"\x00\xff\xff\xff")
            return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
                    + chunk(b"IDAT", raw) + chunk(b"IEND", b""))

        image = self.tmp / "logo.png"
        image.write_bytes(png_1x1())
        prs = Presentation()
        layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = "Alt text deck"
        picture = slide.shapes.add_picture(str(image), Inches(1), Inches(2))
        picture._element.xpath("./*[1]/p:cNvPr")[0].set("descr", "Bar chart: revenue by region, EMEA lowest at 22.8")
        plain = slide.shapes.add_picture(str(image), Inches(3), Inches(2))  # alt text = file name only: ignored
        self.assertEqual(plain._element.xpath("./*[1]/p:cNvPr")[0].get("descr"), "logo.png")

        for holder, text in ((layout, "Layout watermark DRAFT"), (prs.slide_master, "Master legal line 2026")):
            tree = holder.shapes._spTree
            tree.append(CT_Shape.new_textbox_sp(900 + len(tree), "Template text", 0, 0, 100000, 100000))
            shape = [s for s in holder.shapes if s.name == "Template text"][0]
            shape.text_frame.text = text
        path = self.tmp / "alt.pptx"
        prs.save(path)
        return path, layout.name

    def test_pptx_alt_text_is_extracted_and_file_names_are_skipped(self):
        path, _ = self._deck_with_alt_and_template_text()
        alts = [u for u in doc_text.extract_units(path) if u["location"].startswith("slide 1 / alt")]
        self.assertEqual([u["location"] for u in alts], ["slide 1 / alt"])
        self.assertEqual(alts[0]["text"], "Bar chart: revenue by region, EMEA lowest at 22.8")

    def test_pptx_masters_only_with_flag_and_without_prompt_text(self):
        path, layout_name = self._deck_with_alt_and_template_text()
        default = " | ".join(u["text"] for u in doc_text.extract_units(path))
        self.assertNotIn("Layout watermark", default)
        self.assertNotIn("Master legal line", default)
        units = by_location(doc_text.extract_units(path, include_masters=True))
        self.assertEqual(units[f"layout {layout_name}"]["text"], "Layout watermark DRAFT")
        self.assertEqual(units["master"]["text"], "Master legal line 2026")
        joined = " | ".join(u["text"] for u in units.values())
        self.assertNotIn("Click to edit", joined)

    def test_include_masters_cli_flag(self):
        path, layout_name = self._deck_with_alt_and_template_text()
        self.assertNotIn("[master]", h.run_script("doc_text.py", path).stdout)
        self.assertIn("[master] Master legal line 2026", h.run_script("doc_text.py", path, "--include-masters").stdout)


class Cli(unittest.TestCase):
    def test_text_output_format(self):
        proc = h.run_script("doc_text.py", h.WEAK_DECK)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertIn("[slide 4 / title] Sales Overview", proc.stdout)
        self.assertIn("[slide 4 / body] Revenue was $86.2M in Q3\n    This compares to $84.9M in Q2", proc.stdout)

    def test_json_output_is_valid(self):
        proc = h.run_script("doc_text.py", h.PT_DECK, "--json")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        data = json.loads(proc.stdout)
        self.assertEqual(data["format"], "pptx")
        self.assertTrue(all({"location", "kind", "text"} <= set(u) for u in data["units"]))

    def test_titles_flag(self):
        proc = h.run_script("doc_text.py", h.WEAK_DECK, "--titles")
        self.assertEqual(proc.stdout.splitlines()[1], "2. Agenda")

    def test_errors_exit_1(self):
        self.assertEqual(h.run_script("doc_text.py", "/nonexistent/file.pptx").returncode, 1)
        with tempfile.TemporaryDirectory() as tmp:
            bad = Path(tmp) / "x.pdf"
            bad.write_bytes(b"%PDF")
            proc = h.run_script("doc_text.py", bad)
            self.assertEqual(proc.returncode, 1)
            self.assertIn("Unsupported", proc.stderr)

    def test_help_documents_usage(self):
        proc = h.run_script("doc_text.py", "--help")
        self.assertEqual(proc.returncode, 0)
        self.assertIn("--json", proc.stdout)


if __name__ == "__main__":
    unittest.main()
