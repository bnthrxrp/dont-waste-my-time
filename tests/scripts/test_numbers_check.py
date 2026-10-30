"""Tests for numbers_check.py: normalisation, compare (translations) and trace (no invented numbers)."""
import json
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

import _helpers as h
import numbers_check as nc

PT_TEXT = """# Resultados do 3T25

O faturamento foi de R$ 48,7 mi, +6,2% vs. 3T24. A margem caiu -1,8 p.p. para 41,2%.
Foram 1.234 pedidos em 15/11/2026, com prazo até março/27.
Volume de 12 mil caixas; investimento de R$ 1,2 milhão e ticket de R$ 1.234,50.
"""
EN_TEXT = """# Results for Q3 2025

Revenue was BRL 48.7 million, +6.2% vs. Q3 2024. Margin fell -1.8 pp to 41.2%.
There were 1,234 orders on November 15, 2026, with a deadline of March 2027.
Volume of 12,000 boxes; investment of R$1.2M and a ticket of R$ 1,234.50.
"""


def norm(text, locale="en"):
    return [t.norm() for t in nc.tokenize(text, locale)]


def one(text, locale="en"):
    tokens = nc.tokenize(text, locale)
    assert len(tokens) == 1, (text, [t.raw for t in tokens])
    return tokens[0]


class TempFiles(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def write(self, name, text):
        path = self.tmp / name
        path.write_text(text, encoding="utf-8")
        return path

    def compare(self, source_text, target_text, **kwargs):
        src, tgt = self.write("src.md", source_text), self.write("tgt.md", target_text)
        return nc.compare_docs(nc.load_doc(src, kwargs.get("source_locale")), nc.load_doc(tgt, kwargs.get("target_locale")))

    def trace(self, deliverable_text, *source_texts, **kwargs):
        deliverable = self.write("deliverable.md", deliverable_text)
        sources = [self.write(f"source{i}.md", s) for i, s in enumerate(source_texts)]
        return nc.trace_docs(nc.load_doc(deliverable), [nc.load_doc(s) for s in sources], **kwargs)


class Normalisation(unittest.TestCase):
    def test_scale_words_and_currency_are_equivalent(self):
        variants = ["R$ 48,7 mi", "R$ 48,7 milhões", "R$48.7M"]
        self.assertEqual(one("R$ 48,7 mi", "pt").norm(), "BRL 48700000")
        self.assertEqual(one("R$ 48,7 milhões", "pt").norm(), "BRL 48700000")
        self.assertEqual(one("R$48.7M").norm(), "BRL 48700000")
        self.assertEqual(one("BRL 48.7 million").norm(), "BRL 48700000")
        self.assertEqual(one("48,7 milhões de reais", "pt").norm(), "BRL 48700000")
        self.assertEqual(one("48.7 MM").value, Decimal("48700000"))
        self.assertEqual(one("US$ 1.2B").norm(), "USD 1200000000")
        self.assertEqual(one("1,2 bilhão", "pt").value, Decimal("1200000000"))
        self.assertEqual(one("2 tri", "pt").value, Decimal("2000000000000"))
        self.assertEqual(one("12 mil", "pt").value, Decimal("12000"))
        self.assertEqual(one("12k").value, Decimal("12000"))
        self.assertEqual(one("12 thousand").value, Decimal("12000"))
        self.assertEqual(len(variants), 3)

    def test_scale_letters_are_case_sensitive_and_not_greedy(self):
        self.assertEqual(one("5m long").value, Decimal("5"))      # lower-case m is metres, not millions
        self.assertEqual(one("30 min").value, Decimal("30"))      # min is minutes, not "mi" (million)
        self.assertEqual(one("5 kg").value, Decimal("5"))

    def test_percent_points_and_basis_points(self):
        self.assertEqual(one("-1,8 p.p.", "pt").norm(), "-1.8 pp")
        self.assertEqual(one("-1.8 pp").norm(), "-1.8 pp")
        self.assertEqual(one("1.8 percentage points").norm(), "1.8 pp")
        self.assertEqual(one("1,8 pontos percentuais", "pt").norm(), "1.8 pp")
        self.assertEqual(one("50 bps").norm(), "0.5 pp")
        self.assertEqual(one("100 bps").norm(), one("1 p.p.", "pt").norm())
        self.assertEqual(one("6,2 por cento", "pt").norm(), "6.2%")
        self.assertEqual(one("+6,2%", "pt").norm(), one("6.2%").norm())
        self.assertEqual(one("2.5x").norm(), "2.5x")
        self.assertEqual(one("12 pontos de venda", "pt").unit, "")  # points of sale are not percentage points

    def test_dates_normalise_across_languages(self):
        iso = "2026-11-15"
        for text, locale in [("15/11/2026", "pt"), ("November 15, 2026", "en"), ("Nov 15, 2026", "en"),
                             ("15 Nov 2026", "en"), ("2026-11-15", "en"), ("15 de novembro de 2026", "pt"),
                             ("11/15/2026", "en")]:
            token = one(text, locale)
            self.assertEqual((token.kind, token.value), ("date", iso), text)

    def test_ambiguous_numeric_date_follows_document_locale(self):
        self.assertEqual(one("05/11/2026", "pt").value, "2026-11-05")
        self.assertEqual(one("05/11/2026", "en").value, "2026-05-11")

    def test_month_year_and_partial_dates(self):
        self.assertEqual(one("março/27", "pt").value, "2027-03")
        self.assertEqual(one("agosto/26", "pt").value, "2026-08")
        self.assertEqual(one("Outubro de 2026", "pt").value, "2026-10")
        self.assertEqual(one("March 2027").value, "2027-03")
        self.assertEqual(one("Dec 31").value, "--12-31")
        self.assertEqual(one("31 de dezembro", "pt").value, "--12-31")

    def test_ambiguous_month_abbreviations_are_not_dates(self):
        tokens = nc.tokenize("3 out of 5 customers", "en")
        self.assertEqual([(t.kind, t.raw) for t in tokens], [("number", "3"), ("number", "5")])

    def test_quarters_normalise(self):
        for text in ("3T25", "3º trimestre de 2025", "Q3 2025", "3Q25", "Q3'25", "Q3-25", "third quarter of 2025",
                     "terceiro trimestre de 2025"):
            token = one(text, "pt")
            self.assertEqual((token.kind, token.value), ("quarter", "2025-Q3"), text)
        self.assertEqual(one("Q3").value, "Q3")
        self.assertEqual(one("3T", "pt").value, "Q3")

    def test_locale_ambiguity_uses_document_locale(self):
        self.assertEqual(one("1.234", "pt").number, Decimal("1234"))
        self.assertEqual(one("1.234", "en").number, Decimal("1.234"))
        self.assertEqual(one("1,234", "en").number, Decimal("1234"))
        self.assertEqual(one("1,234", "pt").number, Decimal("1.234"))
        # Unambiguous shapes ignore the locale.
        self.assertEqual(one("1.234,5", "en").number, Decimal("1234.5"))
        self.assertEqual(one("1,234.5", "pt").number, Decimal("1234.5"))
        self.assertEqual(one("48,7", "en").number, Decimal("48.7"))
        self.assertEqual(one("6.2", "pt").number, Decimal("6.2"))
        self.assertEqual(one("1.234.567", "en").number, Decimal("1234567"))
        self.assertEqual(one("0.123", "pt").number, Decimal("0.123"))

    def test_document_locale_detection(self):
        pt_units = [{"location": "a", "kind": "line", "text": "Vendas de 1.234 unidades e 48,7 mi"}]
        en_units = [{"location": "a", "kind": "line", "text": "Sales of 1,234 units and 48.7 million"}]
        self.assertEqual(nc.detect_number_locale(pt_units, "pt")[0], "pt")
        self.assertEqual(nc.detect_number_locale(en_units, "en")[0], "en")
        # No evidence either way: fall back to the text language.
        self.assertEqual(nc.detect_number_locale([{"location": "a", "kind": "line", "text": "pedidos 1.234"}], "pt")[0], "pt")

    def test_signs(self):
        self.assertEqual(one("(-1,8 p.p.)", "pt").number, Decimal("-1.8"))
        self.assertEqual(one("−2,5%", "pt").number, Decimal("-2.5"))       # unicode minus
        self.assertEqual(one("+6.2%").number, Decimal("6.2"))
        self.assertEqual(one("R$ -5").number, Decimal("-5"))

    def test_ranges_and_identifiers_are_not_signed_or_counted(self):
        self.assertEqual(norm("4-6 weeks"), ["4", "6"])
        self.assertEqual(norm("2025–2026"), ["2025", "2026"])
        self.assertEqual(norm("SKU REL-24 and B2B and 3PL and COVID-19 and N12"), [])

    def test_references_urls_and_list_numbers_are_ignored(self):
        self.assertEqual(norm("See slide 4 and Table 2, page 7-8."), [])
        self.assertEqual(norm("details at https://example.com/report/2026/12 or a@b.co"), [])
        self.assertEqual(norm("1. First\n2) Second 45%\n### 3. Section 12.5%\n4.2 Pricing 3.3%"), ["45%", "3.3%"])
        self.assertEqual(norm("4.5 Million units\n48.7 MM revenue"), ["4500000", "48700000"])  # scale words are not outline titles

    def test_precision_step(self):
        self.assertEqual(one("48,7 mi", "pt").step, Decimal("100000"))
        self.assertEqual(one("6.25%").step, Decimal("0.01"))
        self.assertEqual(one("41,300").step, Decimal("1"))

    def test_machine_spans_ignore_locale(self):
        text = "Series: 1.234 | 22.8"
        tokens = nc.tokenize(text, "pt", machine_spans=[[8, len(text)]])
        self.assertEqual([t.number for t in tokens], [Decimal("1.234"), Decimal("22.8")])


class Compare(TempFiles):
    def test_correct_translation_has_zero_discrepancies(self):
        result = self.compare(PT_TEXT, EN_TEXT)
        s = result["summary"]
        self.assertEqual((s["missing_in_target"], s["extra_in_target"], s["matched_with_differences"],
                          s["location_count_differences"]), (0, 0, 0, 0), json.dumps(result, indent=1))
        self.assertEqual(s["matched_exact"], 12)
        self.assertEqual(result["source"]["number_format"], "pt")
        self.assertEqual(result["target"]["number_format"], "en")

    def test_altered_number_is_reported_with_a_likely_counterpart(self):
        result = self.compare(PT_TEXT, EN_TEXT.replace("1,234 orders", "1,243 orders"))
        self.assertEqual([m["raw"] for m in result["missing_in_target"]], ["1.234"])
        self.assertEqual([m["raw"] for m in result["extra_in_target"]], ["1,243"])
        self.assertEqual(result["missing_in_target"][0]["likely_counterpart"]["raw"], "1,243")
        self.assertEqual(result["missing_in_target"][0]["location"], "line 4")

    def test_dropped_number_is_reported_missing(self):
        result = self.compare(PT_TEXT, EN_TEXT.replace("a deadline of March 2027", "a deadline"))
        self.assertEqual([m["raw"] for m in result["missing_in_target"]], ["março/27"])
        self.assertEqual(result["extra_in_target"], [])
        self.assertEqual(result["location_count_differences"],
                         [{"location": "line 4", "source_count": 3, "target_count": 2}])

    def test_extra_number_in_target(self):
        result = self.compare(PT_TEXT, EN_TEXT.replace("Volume of 12,000 boxes", "Volume of 12,000 boxes in 7 regions"))
        self.assertEqual([m["raw"] for m in result["extra_in_target"]], ["7"])
        self.assertEqual(result["missing_in_target"], [])

    def test_sign_unit_and_currency_differences(self):
        sign = self.compare("Margem de -1,8 p.p.", "Margin of 1.8 pp")
        self.assertEqual([d["type"] for d in sign["matched_with_differences"]], ["sign"])
        unit = self.compare("Margem de -1,8 p.p.", "Margin of -1.8%")
        self.assertEqual([d["type"] for d in unit["matched_with_differences"]], ["unit"])
        currency = self.compare("Investimento de R$ 1,2 mi", "Investment of 1.2 million")
        self.assertEqual([d["type"] for d in currency["matched_with_differences"]], ["currency"])
        self.assertEqual(currency["summary"]["missing_in_target"], 0)

    def test_plus_sign_does_not_matter(self):
        result = self.compare("Cresceu +6,2%", "Grew 6.2%")
        self.assertEqual(result["summary"]["matched_with_differences"], 0)
        self.assertEqual(result["summary"]["matched_exact"], 1)

    def test_rounding_tolerance_uses_the_coarser_precision(self):
        close = self.compare("Margem de 41,2% e R$ 48,7 mi", "Margin of 41% and BRL 48,683,000")
        self.assertEqual(close["summary"]["missing_in_target"], 0)
        self.assertEqual(close["summary"]["matched_within_rounding"], 2)
        far = self.compare("Margem de 41,2%", "Margin of 41.7%")
        self.assertEqual(far["summary"]["missing_in_target"], 1)

    def test_same_value_twice_needs_two_matches(self):
        result = self.compare("A 5% e B 5%", "A 5%")
        self.assertEqual(result["summary"]["missing_in_target"], 1)

    def test_wrong_locale_format_in_target_is_noted(self):
        text = "Revenue grew 6,2% to R$ 48,7 mi, ahead of the market in the quarter and the year."
        result = self.compare("O faturamento cresceu 6,2% para R$ 48,7 mi, acima do mercado no trimestre e no ano.", text)
        self.assertEqual(result["summary"]["missing_in_target"], 0)
        self.assertTrue(any("pt-BR format" in note for note in result["notes"]), result["notes"])

    def test_locale_overrides(self):
        wrong = self.compare("Foram 1.234 pedidos", "There were 1.234 orders", source_locale="en", target_locale="en")
        self.assertEqual(wrong["summary"]["missing_in_target"], 0)       # both read as 1.234
        right = self.compare("Foram 1.234 pedidos", "There were 1,234 orders", source_locale="pt", target_locale="en")
        self.assertEqual(right["summary"]["missing_in_target"], 0)       # both read as 1234
        mismatch = self.compare("Foram 1.234 pedidos", "There were 1.234 orders", source_locale="pt", target_locale="en")
        self.assertEqual(mismatch["summary"]["missing_in_target"], 1)    # 1234 vs 1.234: a real error

    def test_pptx_deck_compared_with_itself_has_no_discrepancies(self):
        result = nc.compare_docs(nc.load_doc(h.PT_DECK), nc.load_doc(h.PT_DECK))
        s = result["summary"]
        self.assertEqual((s["missing_in_target"], s["extra_in_target"], s["location_count_differences"]), (0, 0, 0))
        self.assertGreater(s["matched_exact"], 50)

    def test_pptx_translation_with_one_changed_cell(self):
        from pptx import Presentation
        prs = Presentation(str(h.PT_DECK))
        cell = next(sh for sh in prs.slides[2].shapes if getattr(sh, "has_table", False) and sh.has_table).table.cell(1, 1)
        self.assertEqual(cell.text, "21,4")
        cell.text_frame.paragraphs[0].runs[0].text = "21,5"
        changed = self.tmp / "changed.pptx"
        prs.save(changed)
        result = nc.compare_docs(nc.load_doc(h.PT_DECK), nc.load_doc(changed))
        self.assertEqual([(m["raw"], m["location"]) for m in result["missing_in_target"]], [("21,4", "slide 3 / table r2c2")])
        self.assertEqual(result["extra_in_target"][0]["raw"], "21,5")

    def test_chart_values_are_machine_numbers_in_a_portuguese_deck(self):
        doc = nc.load_doc(h.PT_DECK)
        chart = [t for t in doc.tokens if t.location == "slide 4 / chart"]
        self.assertIn(Decimal("74.8"), [t.number for t in chart if t.kind == "number"])
        self.assertIn("2026-Q3", [t.value for t in chart if t.kind == "quarter"])


class Trace(TempFiles):
    SOURCE = "Revenue was $86.2M in Q3; last year $88.7M. Gross margin 23.4% vs 27.9%. Volume 41,300 tonnes."

    def untraced(self, deliverable, *sources, **kw):
        return [u["raw"] for u in self.trace(deliverable, *(sources or (self.SOURCE,)), **kw)["untraced"]]

    def test_invented_number_is_flagged(self):
        self.assertEqual(self.untraced("Revenue was $86.2M and we expect 17% growth."), ["17%"])

    def test_rounded_source_number_is_not_flagged(self):
        self.assertEqual(self.untraced("Revenue was about $86M, margin 23%, volume 41k tonnes."), [])
        self.assertEqual(self.untraced("Revenue was $86,200,000."), [])

    def test_scale_and_unit_differences_are_tolerated_but_listed(self):
        result = self.trace("Revenue was R$ 86.2M", "Revenue (R$ millions): 86.2")
        self.assertEqual(result["untraced"], [])
        self.assertEqual(result["summary"]["traced_loosely"], 1)

    def test_derived_number_is_untraced_and_gets_a_basis_hint(self):
        result = self.trace("Revenue fell 2.8% from $88.7M to $86.2M.", self.SOURCE)
        self.assertEqual([u["raw"] for u in result["untraced"]], ["2.8%"])
        self.assertIn("percent change", result["untraced"][0]["possible_basis"])
        self.assertIn("derived", " ".join(result["notes"]))

    def test_years_ignored_by_default_and_checked_on_request(self):
        text = "In 2027 we plan to grow. Q4 2027 and March 2027 matter. 86.2 is fine."
        self.assertEqual(self.untraced(text), [])
        self.assertEqual(self.untraced(text, ignore_years=True), [])
        raw = [u["raw"] for u in self.trace(text, self.SOURCE, ignore_years=False)["untraced"]]
        self.assertEqual(sorted(raw), ["2027", "March 2027", "Q4 2027"])

    def test_full_dates_are_always_checked(self):
        self.assertEqual(self.untraced("Decision by December 31, 2026."), ["December 31, 2026"])
        self.assertEqual(self.untraced("Decision by Dec 31.", "Notice must be given by Dec 31."), [])

    def test_numbering_slide_refs_and_urls_are_ignored(self):
        text = "1. Revenue was $86.2M.\n2) See slide 12 and https://x.example/a/987.\n### 3. Outlook"
        self.assertEqual(self.untraced(text), [])

    def test_grouped_by_location_with_snippets(self):
        result = self.trace("Revenue was $86.2M.\n\nOther: we added 31 stores and 7 plants.", self.SOURCE)
        locations = result["untraced_by_location"]
        self.assertEqual(list(locations), ["line 3"])
        self.assertEqual([i["raw"] for i in locations["line 3"]], ["31", "7"])
        self.assertIn("31 stores", locations["line 3"][0]["snippet"])

    def test_multiple_sources_including_markdown_log_csv_and_pptx(self):
        log = self.write("log.md", "# Source log\n- Brandt contract loses about $1.4M per year (finance, 2026-10-02)\n")
        csv = self.write("regions.csv", "region,rev\nEMEA,22.8\nNA,31.4\n")
        deliverable = self.write("d.md", "Brandt loses $1.4M per year. EMEA revenue was $22.8M and NA $31.4M. "
                                         "Plant 2 runs at 64%, up 9%.\n")
        result = nc.trace_docs(nc.load_doc(deliverable),
                               [nc.load_doc(log), nc.load_doc(csv), nc.load_doc(h.WEAK_DECK)])
        self.assertEqual([u["raw"] for u in result["untraced"]], ["9%"])  # 64% is on slide 8 of the weak deck
        self.assertEqual(len(result["sources"]), 3)

    def test_one_source_number_can_back_many_deliverable_numbers(self):
        self.assertEqual(self.untraced("Revenue $86.2M. Again, $86.2M. And $86.2M."), [])


class Cli(TempFiles):
    def test_compare_cli_text_and_exit_codes(self):
        src, tgt = self.write("pt.md", PT_TEXT), self.write("en.md", EN_TEXT)
        ok = h.run_script("numbers_check.py", "compare", src, tgt)
        self.assertEqual(ok.returncode, 0, ok.stderr)
        self.assertIn("RESULT: OK", ok.stdout)
        bad_target = self.write("bad.md", EN_TEXT.replace("1,234 orders", "1,243 orders"))
        bad = h.run_script("numbers_check.py", "compare", src, bad_target)
        self.assertEqual(bad.returncode, 0)  # discrepancies are findings, not errors
        self.assertIn("(a) Source numbers missing from target (1)", bad.stdout)
        self.assertIn("likely edited to '1,243'", bad.stdout)
        strict = h.run_script("numbers_check.py", "compare", src, bad_target, "--strict")
        self.assertEqual(strict.returncode, 2)

    def test_compare_cli_json(self):
        src, tgt = self.write("pt.md", PT_TEXT), self.write("en.md", EN_TEXT)
        data = json.loads(h.run_script("numbers_check.py", "compare", src, tgt, "--json").stdout)
        self.assertEqual(data["command"], "compare")
        self.assertEqual(data["summary"]["missing_in_target"], 0)

    def test_trace_cli(self):
        source = self.write("s.md", "Revenue $86.2M, margin 23.4%.")
        deliverable = self.write("d.md", "Revenue $86.2M, margin 23.4%, forecast 15%.")
        proc = h.run_script("numbers_check.py", "trace", deliverable, "--sources", source)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertIn("15%", proc.stdout)
        self.assertIn("derived calculations", proc.stdout)
        self.assertEqual(h.run_script("numbers_check.py", "trace", deliverable, "--sources", source, "--strict").returncode, 2)
        clean = self.write("c.md", "Revenue $86.2M.")
        self.assertEqual(h.run_script("numbers_check.py", "trace", clean, "--sources", source, "--strict").returncode, 0)
        data = json.loads(h.run_script("numbers_check.py", "trace", deliverable, "--sources", source, "--json").stdout)
        self.assertEqual(data["summary"]["untraced"], 1)

    def test_include_years_flag(self):
        source = self.write("s.md", "Revenue $86.2M.")
        deliverable = self.write("d.md", "Revenue $86.2M in 2031.")
        default = json.loads(h.run_script("numbers_check.py", "trace", deliverable, "--sources", source, "--json").stdout)
        self.assertEqual(default["summary"]["untraced"], 0)
        years = json.loads(h.run_script("numbers_check.py", "trace", deliverable, "--sources", source,
                                        "--include-years", "--json").stdout)
        self.assertEqual(years["summary"]["untraced"], 1)

    def test_io_errors_exit_1(self):
        self.assertEqual(h.run_script("numbers_check.py", "compare", "/nope.md", "/nope2.md").returncode, 1)
        self.assertEqual(h.run_script("numbers_check.py", "trace", "/nope.md", "--sources", "/nope2.md").returncode, 1)

    def test_help(self):
        out = h.run_script("numbers_check.py", "--help").stdout
        self.assertIn("compare", out)
        self.assertIn("trace", out)
        self.assertIn("--source-locale", h.run_script("numbers_check.py", "compare", "--help").stdout)


if __name__ == "__main__":
    unittest.main()
